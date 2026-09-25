"""Stage 2 of the pipeline: load every source into DuckDB as text, and log row counts.

Loading as text first means a badly formatted value never breaks the load:
it arrives exactly as the source sent it, and the staging layer decides how to
type it, which is where every cleaning rule is written down.
"""

import time

import pandas as pd
from openpyxl import load_workbook

from src.config import RAW_DIR, REFERENCE_DIR, WAREHOUSE, connect
from src.generate.writers import SOURCE_FILES

SCHEMAS = ["raw", "reference", "staging", "marts", "kpi", "dq"]

# (table in the raw schema, source system, format)
SOURCES = [
    ("app_events", "Mobile app analytics", "json"),
    ("customers", "Core banking", "csv"),
    ("accounts", "Core banking", "csv"),
    ("account_balances", "Core banking", "csv"),
    ("ledger_postings", "Posting engine", "parquet"),
    ("card_authorisations", "Card processor", "csv"),
    ("loans", "Loan management system", "parquet"),
    ("loan_tape", "Loan management system", "parquet"),
    ("marketing_spend", "Marketing team spreadsheet", "excel"),
]

APP_EVENT_COLUMNS = ["event_id", "event_name", "event_ts", "received_at", "applicant_id", "context", "properties"]

# Reference columns that must stay text (codes with leading zeros, labels that look like numbers)
REFERENCE_TEXT_COLUMNS = {
    "response_code_map": ["response_code"],
    "product_map": ["raw_code", "tenor_months"],
    "products": ["tenor_months"],
    "channel_map": ["raw_label"],
}


def _read_marketing_sheet(path):
    """Read the hand-kept sheet as text, starting from its header row."""
    ws = load_workbook(path, data_only=False).active
    rows = list(ws.iter_rows(values_only=True))
    header_at = next(i for i, r in enumerate(rows) if r and str(r[0]).strip().lower() == "channel")
    header = [str(v) if v is not None else f"column_{j + 1}" for j, v in enumerate(rows[header_at])]
    body = []
    for r in rows[header_at + 1:]:
        if all(v is None for v in r):
            continue
        body.append([None if v is None else str(v) for v in r])
    df = pd.DataFrame(body, columns=header, dtype=object)
    df.insert(0, "sheet_row", range(1, len(df) + 1))
    df["sheet_row"] = df["sheet_row"].astype(str)
    return df


def load_source(con, table, fmt):
    path = RAW_DIR / SOURCE_FILES[table]
    target = f"raw.{table}"
    if fmt == "csv":
        con.execute(f"CREATE TABLE {target} AS SELECT * FROM read_csv('{path.as_posix()}', all_varchar = true, header = true)")
    elif fmt == "parquet":
        con.execute(f"CREATE TABLE {target} AS SELECT COLUMNS(*)::VARCHAR FROM read_parquet('{path.as_posix()}')")
    elif fmt == "json":
        cols = ", ".join(f"'{c}': 'VARCHAR'" for c in APP_EVENT_COLUMNS)
        con.execute(f"CREATE TABLE {target} AS SELECT * FROM read_json('{path.as_posix()}', "
                    f"format = 'newline_delimited', columns = {{{cols}}})")
    elif fmt == "excel":
        df = _read_marketing_sheet(path)
        con.register("sheet_df", df)
        con.execute(f"CREATE TABLE {target} AS SELECT * FROM sheet_df")
        con.unregister("sheet_df")
    return con.execute(f"SELECT count(*) FROM {target}").fetchone()[0]


def load_reference(con):
    loaded = 0
    for f in sorted(REFERENCE_DIR.glob("*.csv")):
        name = f.stem
        text_cols = REFERENCE_TEXT_COLUMNS.get(name, [])
        types = ", ".join(f"'{c}': 'VARCHAR'" for c in text_cols)
        extra = f", types = {{{types}}}" if text_cols else ""
        con.execute(f"CREATE TABLE reference.{name} AS SELECT * FROM read_csv('{f.as_posix()}', header = true{extra})")
        loaded += 1
    return loaded


def main():
    started = time.perf_counter()
    missing = [SOURCE_FILES[t] for t, _, _ in SOURCES if not (RAW_DIR / SOURCE_FILES[t]).exists()]
    if missing:
        raise SystemExit(f"Missing source files in data/raw/: {', '.join(missing)}. Run the generate stage first.")
    WAREHOUSE.unlink(missing_ok=True)
    WAREHOUSE.with_suffix(".duckdb.wal").unlink(missing_ok=True)
    con = connect()
    for schema in SCHEMAS:
        con.execute(f"CREATE SCHEMA {schema}")

    log = []
    for table, system, fmt in SOURCES:
        rows = load_source(con, table, fmt)
        log.append((table, system, SOURCE_FILES[table], fmt, rows))
        print(f"      raw.{table:<22} {rows:>10,} rows  ({fmt}, {system})")
    con.execute("CREATE TABLE raw._load_log (source_table VARCHAR, source_system VARCHAR, file_name VARCHAR, "
                "file_format VARCHAR, row_count BIGINT)")
    con.executemany("INSERT INTO raw._load_log VALUES (?, ?, ?, ?, ?)", log)
    n_ref = load_reference(con)
    con.close()
    print(f"      {n_ref} reference tables loaded in {time.perf_counter() - started:.1f} s")


if __name__ == "__main__":
    main()
