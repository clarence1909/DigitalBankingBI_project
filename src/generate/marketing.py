"""The marketing team's hand-kept spend tracker (source 6, Excel).

It is deliberately laid out the way a person maintains it: a title block,
months across the columns, numbers typed as text in places, a few header
cells that Excel turned into real dates, a notes column, and a TOTAL row
typed by hand that does not always foot.
"""

from datetime import datetime

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.util import normalise_xlsx

from . import params as P
from .timeutil import MONTHS, MONTH_LABELS, N_MONTHS

SHEET_LINES = [
    # (label typed in the sheet, channel it pays for or None, note)
    ("Google Ads", "google_search", "Search campaigns, billed monthly"),
    ("Facebook / Instagram", "meta_ads", "Meta Ads Manager invoices"),
    ("TikTok", "tiktok", "Includes agency fee from Oct-25"),
    ("Affiliate (cashback)", "affiliate", "RM18 per account opened + RM1,500 platform fee"),
    ("Referral programme", "referral", "RM25 to referrer + RM10 to friend, per account opened"),
    ("Organic / App Store", "organic", "No paid spend"),
    ("FD Raya promo (in-app banners)", None, "Deposit campaign, not acquisition"),
    ("Brand / PR", None, "Launch event, festive brand ads"),
]
BRAND_SPEND = {"2024-09": 25000, "2025-01": 6000, "2025-03": 8000, "2026-01": 6500, "2026-03": 7500}
PROMO_SPEND = {"2025-03": 4800, "2025-04": 5200, "2025-05": 3900}
DATE_HEADER_MONTHS = {"2025-01", "2025-07", "2026-01"}   # Excel auto-converted these to dates
TOTAL_TYPO = ("2026-01", 500.0)   # the hand-typed TOTAL is RM500 too high this month
MISSING_INVOICE = ("Facebook / Instagram", "2026-08")   # typed as n/a: invoice not in yet


def spend_table(rng, starts, opened):
    """Spend in RM per sheet line per month. starts/opened: DataFrames month x channel."""
    table = {}
    for label, ch, _ in SHEET_LINES:
        row = np.zeros(N_MONTHS)
        if ch in P.COST_PER_START:
            row = starts[ch].to_numpy() * P.COST_PER_START[ch] * rng.lognormal(0, 0.03, N_MONTHS)
        elif ch == "affiliate":
            row = opened[ch].to_numpy() * P.COST_PER_OPENED[ch] + P.AFFILIATE_PLATFORM_FEE
        elif ch == "referral":
            row = opened[ch].to_numpy() * P.COST_PER_OPENED[ch]
        elif label.startswith("FD Raya"):
            row = np.array([PROMO_SPEND.get(m, 0.0) for m in MONTH_LABELS])
        elif label.startswith("Brand"):
            row = np.array([BRAND_SPEND.get(m, 0.0) for m in MONTH_LABELS])
        table[label] = np.round(row, 2)
    return table


def write_sheet(rng, table, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Spend 2024-26"
    ws["A1"] = "Kelip Bank - Marketing spend tracker"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Owner: Growth team. All figures in RM. Updated by hand at month end."
    ws["A2"].font = Font(italic=True, color="666666")

    header_row = 4
    ws.cell(header_row, 1, "Channel")
    for j, (label, period) in enumerate(zip(MONTH_LABELS, MONTHS), start=2):
        if label in DATE_HEADER_MONTHS:
            cell = ws.cell(header_row, j, datetime(period.year, period.month, 1))
            cell.number_format = "mmm-yy"
        else:
            ws.cell(header_row, j, period.strftime("%b-%y"))
    total_col = N_MONTHS + 2
    ws.cell(header_row, total_col, "Total")
    ws.cell(header_row, total_col + 1, "Notes")
    for c in range(1, total_col + 2):
        ws.cell(header_row, c).font = Font(bold=True)
        ws.cell(header_row, c).fill = PatternFill("solid", fgColor="DDEBF7")

    r = header_row + 1
    first_data_row = r
    for label, ch, note in SHEET_LINES:
        ws.cell(r, 1, label)
        values = table[label]
        for j, v in enumerate(values, start=2):
            if ch == "organic":
                ws.cell(r, j, "-")
            elif (label, MONTH_LABELS[j - 2]) == MISSING_INVOICE:
                ws.cell(r, j, "n/a")
            elif v == 0:
                continue   # left blank
            else:
                u = rng.random()
                if u < 0.06:
                    ws.cell(r, j, f"{v:,.2f}")            # typed with thousands separators
                elif u < 0.09:
                    ws.cell(r, j, f"RM {v:,.2f}")         # typed with a currency prefix
                else:
                    cell = ws.cell(r, j, float(v))
                    cell.number_format = "#,##0.00"
        ws.cell(r, total_col, f"=SUM(B{r}:{get_column_letter(total_col - 1)}{r})")
        ws.cell(r, total_col + 1, note)
        r += 1
    last_data_row = r - 1
    r += 1   # blank row, then the hand-typed TOTAL
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    for j, label in enumerate(MONTH_LABELS, start=2):
        total = sum(table[l][j - 2] for l, ch, _ in SHEET_LINES
                    if ch != "organic" and (l, label) != MISSING_INVOICE)
        if label == TOTAL_TYPO[0]:
            total += TOTAL_TYPO[1]
        cell = ws.cell(r, j, round(float(total), 2))
        cell.number_format = "#,##0.00"
        cell.font = Font(bold=True)
    ws.cell(r + 2, 1, "n/a = invoice not received yet").font = Font(italic=True, color="666666")

    ws.column_dimensions["A"].width = 32
    for j in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(j)].width = 12
    ws.column_dimensions[get_column_letter(total_col + 1)].width = 48
    ws.freeze_panes = ws.cell(header_row + 1, 2)
    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    wb.properties.creator = "Growth team"
    wb.properties.created = datetime(2024, 9, 1)
    wb.properties.modified = datetime(2026, 9, 1)
    wb.save(path)
    normalise_xlsx(path)
