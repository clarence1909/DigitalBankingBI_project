"""Recalculate the Excel KPI pack in LibreOffice and check it against the warehouse.

openpyxl writes formulas without their results, so on its own the file proves
nothing about the numbers. This check opens a copy in LibreOffice (headless),
recalculates every formula, saves it, then reads the results back and tests:

  1. no cell holds an error (#REF!, #NAME?, #DIV/0!, #N/A, ...)
  2. the report sheets hold formulas, not pasted numbers (outside the input cells)
  3. every value on Monthly KPIs matches kpi.kpi_monthly
  4. the Scorecard matches kpi.scorecard for the month it opens on
  5. the Trend sheet's actual, plan and status match the warehouse for every month
  6. every row on the Checks sheet says PASS

The results go to reports/excel_pack_check.md. If LibreOffice is not
installed the check is skipped with a message (CI installs it, so it always
runs there).
"""

import math
import os
import platform
import shutil
import subprocess
import tempfile
import warnings
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.config import REPORTS_DIR

from .excel_pack import CHECKS_FIRST_ROW, CONTROLS, PACK, REPORT_MONTH, TREND_FIRST_ROW

REPORT = REPORTS_DIR / "excel_pack_check.md"
ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")
REL_TOL = 1e-9
SOFFICE = [
    "soffice", "libreoffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
MACRO_URL = "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application"

# Cells on the report sheets that may hold typed-in numbers or dates
INPUT_CELLS = {
    "Scorecard": lambda r, c: (r, c) == (5, 4),                      # the reporting month
    "Trend": lambda r, c: c == 2 and r >= TREND_FIRST_ROW,           # the month labels
    "Monthly KPIs": lambda r, c: r == 6,                             # the month headers
    "Checks": lambda r, c: c in (4, 6),                              # SQL control totals, tolerances
}


def find_soffice():
    for candidate in SOFFICE:
        found = shutil.which(candidate) or (candidate if Path(candidate).is_file() else None)
        if found:
            return found
    return None


def recalculate(soffice, source, workdir):
    """Copy the pack into workdir, recalculate it in LibreOffice and return the copy's path."""
    target = workdir / source.name
    shutil.copyfile(source, target)
    profile = workdir / "profile"
    profile_url = profile.as_uri()
    env = os.environ.copy()
    if platform.system() != "Windows":
        env.setdefault("SAL_USE_VCLPLUGIN", "svp")
    # First start creates a user profile; the macro then goes into its Standard library
    subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={profile_url}"],
                   env=env, capture_output=True, timeout=180)
    macro_dir = profile / "user" / "basic" / "Standard"
    if not macro_dir.is_dir():
        raise RuntimeError("LibreOffice did not create a user profile, so the pack was not recalculated")
    (macro_dir / "Module1.xba").write_text(MACRO, encoding="utf-8")
    before = target.stat().st_mtime_ns, target.stat().st_size
    result = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={profile_url}",
                             MACRO_URL, str(target)], env=env, capture_output=True, text=True, timeout=300)
    if result.returncode != 0 or (target.stat().st_mtime_ns, target.stat().st_size) == before:
        raise RuntimeError(f"LibreOffice did not recalculate the pack: {result.stderr.strip() or result.returncode}")
    return target


def close(a, b):
    if a is None or b is None or a == "" or b == "":
        return (a in (None, "")) and (b in (None, ""))
    return math.isclose(float(a), float(b), rel_tol=REL_TOL, abs_tol=1e-9)


def error_cells(values):
    found = []
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in ERRORS:
                    found.append(f"{ws.title}!{cell.coordinate} {cell.value}")
    return found


def is_number(value):
    return isinstance(value, (int, float, date)) and not isinstance(value, bool)


def formula_census(formulas):
    """Count formulas, and list typed-in numbers on the report sheets outside the input cells."""
    count, pasted = 0, []
    for ws in formulas.worksheets:
        allowed = INPUT_CELLS.get(ws.title)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
                elif allowed is not None and is_number(cell.value) and not allowed(cell.row, cell.column):
                    pasted.append(f"{ws.title}!{cell.coordinate}")
    return count, pasted


def check_monthly(ws, kpi):
    """Every cell on Monthly KPIs against kpi.kpi_monthly."""
    expected = {(r.kpi_id, r.month_start.date()): r.value for r in kpi.itertuples()}
    months = {c: ws.cell(row=6, column=c).value for c in range(5, 5 + 24)}
    compared, problems = 0, []
    for r in range(7, ws.max_row + 1):
        kpi_id = ws.cell(row=r, column=2).value
        if not kpi_id:
            continue
        for c, month in months.items():
            got = ws.cell(row=r, column=c).value
            want = expected.get((kpi_id, month.date()))
            compared += 1
            if not close(got, want):
                problems.append(f"Monthly KPIs {kpi_id} {month:%Y-%m}: pack {got}, warehouse {want}")
    return compared, problems


def check_scorecard(ws, scorecard):
    sc = scorecard[scorecard["month_start"].dt.date == REPORT_MONTH].set_index("kpi_id")
    compared, problems = 0, []
    for r in range(10, 10 + len(sc)):
        kpi_id = ws.cell(row=r, column=2).value
        want = sc.loc[kpi_id]
        pairs = {
            "actual": (ws.cell(row=r, column=5).value, want["actual"]),
            "plan": (ws.cell(row=r, column=6).value, want["plan"]),
            "% of plan": (ws.cell(row=r, column=7).value, want["pct_of_plan"]),
            "variance": (ws.cell(row=r, column=8).value, want["variance_to_plan"]),
            "prior month": (ws.cell(row=r, column=11).value, want["prior_month"]),
        }
        for name, (got, expected) in pairs.items():
            compared += 1
            if not close(got, expected):
                problems.append(f"Scorecard {kpi_id} {name}: pack {got}, warehouse {expected}")
        for name, got, expected in [("status", ws.cell(row=r, column=10).value, want["status_label"]),
                                    ("change", ws.cell(row=r, column=12).value, want["change_vs_prior_month"])]:
            compared += 1
            if str(got).lower() != str(expected).lower():
                problems.append(f"Scorecard {kpi_id} {name}: pack {got}, warehouse {expected}")
    return compared, problems


def check_trend(ws, kpi, scorecard):
    divisor = ws["C10"].value
    kpi_id = ws["C6"].value
    actual = {r.month_start.date(): r.value for r in kpi[kpi["kpi_id"] == kpi_id].itertuples()}
    sc = scorecard[scorecard["kpi_id"] == kpi_id]
    plan = {r.month_start.date(): (r.plan, r.status_label) for r in sc.itertuples()}
    compared, problems = 0, []
    for r in range(TREND_FIRST_ROW, ws.max_row + 1):
        month = ws.cell(row=r, column=2).value
        if not hasattr(month, "year"):
            break
        month = month.date()
        got_actual, got_plan = ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value
        want_plan, want_status = plan.get(month, (None, None))
        checks = [("actual", got_actual * divisor if is_number(got_actual) else None, actual.get(month)),
                  ("plan", got_plan * divisor if is_number(got_plan) else None, want_plan)]
        for name, got, expected in checks:
            compared += 1
            if not close(got, expected):
                problems.append(f"Trend {kpi_id} {month:%Y-%m} {name}: pack {got}, warehouse {expected}")
        compared += 1
        if str(ws.cell(row=r, column=7).value).lower() != str(want_status).lower():
            problems.append(f"Trend {kpi_id} {month:%Y-%m} status: pack {ws.cell(row=r, column=7).value}, "
                            f"warehouse {want_status}")
    return compared, problems, kpi_id


def check_controls(ws):
    rows, problems = [], []
    for i, control in enumerate(CONTROLS):
        r = CHECKS_FIRST_ROW + i
        label, pack, sql, diff, result = (ws.cell(row=r, column=c).value for c in (2, 3, 4, 5, 7))
        rows.append((label, pack, sql, diff, result))
        if result != "PASS":
            problems.append(f"Checks: '{label}' says {result} (pack {pack}, SQL {sql})")
    return rows, problems


def write_report(summary, controls, problems):
    lines = [
        "# Excel KPI pack check",
        "",
        "Generated by `src/export/excel_check.py`. The pack (`reports/kelip_bank_kpi_pack.xlsx`) is recalculated "
        "in LibreOffice, then every figure below is compared with the warehouse. All data is synthetic.",
        "",
        "| Test | Result |",
        "|---|---|",
    ]
    lines += [f"| {name} | {result} |" for name, result in summary]
    lines += ["", "## Control totals on the Checks sheet", "",
              "| Check | In the pack | Control total from SQL | Difference | Result |", "|---|---:|---:|---:|---|"]
    for (label, pack, sql, diff, result), control in zip(controls, CONTROLS):
        fmt = control[4]
        show = (lambda v: f"{v:,.0f}") if fmt == "#,##0" else (lambda v: f"{v:.6%}") if "%" in fmt \
            else (lambda v: f"RM{v:,.2f}")
        lines.append(f"| {label} | {show(pack)} | {show(sql)} | {diff:.1e} | {result} |")
    if problems:
        lines += ["", "## Problems", ""] + [f"- {p}" for p in problems[:50]]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run(con):
    soffice = find_soffice()
    if soffice is None:
        print("      LibreOffice not found, so the Excel recalculation check was skipped "
              "(install LibreOffice to run it; CI always does)")
        return None
    kpi = con.execute("SELECT kpi_id, month_start, value FROM kpi.kpi_monthly").df()
    kpi["month_start"] = pd.to_datetime(kpi["month_start"])
    scorecard = con.execute("SELECT * FROM kpi.scorecard").df()
    scorecard["month_start"] = pd.to_datetime(scorecard["month_start"])

    with tempfile.TemporaryDirectory(prefix="kelip_excel_check_") as tmp:
        recalculated = recalculate(soffice, PACK, Path(tmp))
        with warnings.catch_warnings():   # openpyxl cannot parse LibreOffice's chart XML; the chart is not checked
            warnings.filterwarnings("ignore", message="Unable to read chart", category=UserWarning)
            values = load_workbook(recalculated, data_only=True)
        formulas = load_workbook(PACK)
        errors = error_cells(values)
        formula_count, pasted = formula_census(formulas)
        monthly_n, monthly_problems = check_monthly(values["Monthly KPIs"], kpi)
        score_n, score_problems = check_scorecard(values["Scorecard"], scorecard)
        trend_n, trend_problems, trend_kpi = check_trend(values["Trend"], kpi, scorecard)
        controls, control_problems = check_controls(values["Checks"])

    problems = (errors + [f"Pasted number on a report sheet: {p}" for p in pasted] + monthly_problems
                + score_problems + trend_problems + control_problems)
    passes = sum(1 for *_, result in controls if result == "PASS")
    summary = [
        ("Formulas recalculated", f"{formula_count:,}"),
        ("Cells with an error value", f"{len(errors)}"),
        ("Typed-in numbers on report sheets, outside the input cells", f"{len(pasted)}"),
        ("Monthly KPIs cells matching kpi.kpi_monthly", f"{monthly_n - len(monthly_problems):,} of {monthly_n:,}"),
        (f"Scorecard values matching kpi.scorecard ({REPORT_MONTH:%b %Y})",
         f"{score_n - len(score_problems)} of {score_n}"),
        (f"Trend values matching the warehouse ({trend_kpi})", f"{trend_n - len(trend_problems)} of {trend_n}"),
        ("Checks sheet rows that say PASS", f"{passes} of {len(controls)}"),
    ]
    write_report(summary, controls, problems)
    if problems:
        for p in problems[:20]:
            print(f"      {p}")
        raise SystemExit(f"The Excel KPI pack failed its check: {len(problems)} problem(s); "
                         f"see {REPORT.relative_to(REPORTS_DIR.parent)}")
    print(f"      Excel pack: {formula_count:,} formulas, 0 errors, "
          f"{monthly_n + score_n + trend_n:,} values match the warehouse, {passes} of {len(controls)} checks pass")
    return summary
